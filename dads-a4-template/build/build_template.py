# -*- coding: utf-8 -*-
"""
DADS A4タテ PowerPointテンプレート ビルダー

  python build_template.py [出力パス]

デジタル庁「ダッシュボードイメージ作成ツールキット」(16:9) のレイアウト思想を、
デジタル庁デザインシステム(DADS)の基本デザインに従ってA4タテへ再設計する。

出典：デジタル庁デザインシステムウェブサイト https://design.digital.go.jp/dads/
"""
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

from decimal import Decimal, ROUND_HALF_UP

import spec
from oxml_helpers import (emu, ph_sp, rect_sp, text_sp, sldnum_sp, set_bg,
                          clear_shapes, para, rpr)

# 判型・グリッド・タイポグラフィは spec のプロファイルから解決する（DADS-002）。
# 環境変数 DADS_PROFILE で切り替える。1プロセス=1プロファイル。
S = spec.load(os.environ.get("DADS_PROFILE", spec.DEFAULT_PROFILE))

# A4の版面幅。サンプルページ内の固定寸法を判型に合わせて比例させる基準に使う
A4_CONTENT_W = spec.load("a4").CONTENT_W

def dads_token(st):
    """テキストスタイルから DADS のトークン名を組み立てる。

    判型によって実サイズが変わるため、固定文字列で書くと誤った仕様を表示してしまう。
    DADSの系統は Display(48px以上) / Dense(14px以下) / Standard(その間)。
    """
    px, lh = st["px"], st["lh"]
    fam = "Dsp" if px >= 48 else ("Dns" if px <= 14 else "Std")
    return f'{fam}-{px}{"B" if st["b"] else "N"}-{lh}'


def dads_defined(st):
    """DADS のテキストスタイル表に実在する組み合わせか。

    b5-compact の本文（14px・行高170%）は表に無い組み合わせであり、
    そのことをテンプレート上でも示す必要がある（設計書 DADS-002 §3）。
    """
    px, lh = st["px"], st["lh"]
    if px >= 48:
        return lh == 140
    if px <= 14:
        return lh in (130, 120, 100)
    return lh in (140, 150, 160, 170, 175)


# zip に書かれる更新日時。1980-01-01 00:00:00 は zip 形式が表現できる最小値で、
# 再現可能ビルドの慣習として広く使われている（同じ内容なら同じバイト列になる）。
# 生成時刻を入れると、中身が同じでも毎回 git の差分になり、本当の変更を見分けられなくなる。
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def freeze_zip_dates(path):
    """pptx(zip) 内の更新日時を固定し、同じ内容なら同じバイト列になるようにする。

    パートの順序・圧縮方式・属性はそのまま保つ。順序が変わると PowerPoint が
    読めなくなることがあるため、infolist() の順に書き直す。
    """
    src = zipfile.ZipFile(path)
    items = [(i, src.read(i.filename)) for i in src.infolist()]
    src.close()
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in items:
            ni = zipfile.ZipInfo(info.filename, date_time=ZIP_EPOCH)
            ni.compress_type = info.compress_type
            ni.external_attr = info.external_attr
            ni.internal_attr = info.internal_attr
            ni.create_system = info.create_system
            out.writestr(ni, data)
    os.replace(tmp, path)


def mm1(v):
    """mm値をサンプル本文の表記（小数第1位・四捨五入）へ整える。

    座標計算には使わない。表示専用。
    浮動小数点誤差を先に落とす（U*3 は 6.349999999999999 になり、
    そのまま丸めると 6.3 になってしまう。正しくは 6.35 -> 6.4）。
    """
    return str(Decimal(repr(round(v, 9))).quantize(Decimal("0.1"),
                                                   rounding=ROUND_HALF_UP))


OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "dist", "DADS_A4_Portrait_Template.pptx")

FONT = "Noto Sans JP"
FONT_FALLBACK = "Yu Gothic Medium"   # Noto Sans JP 未導入環境での代替


# ============================================================ テーマ
def rewrite_theme(prs):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")

    clr = (
        f'<a:clrScheme name="デジタル庁デザインシステム {S.PAPER_NAME.replace("タテ", "")}">'
        f'<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
        f'<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
        f'<a:dk2><a:srgbClr val="{S.PRIMARY}"/></a:dk2>'
        f'<a:lt2><a:srgbClr val="{S.BG_TINT}"/></a:lt2>'
        f'<a:accent1><a:srgbClr val="{S.PRIMARY}"/></a:accent1>'
        f'<a:accent2><a:srgbClr val="{S.SECONDARY}"/></a:accent2>'
        f'<a:accent3><a:srgbClr val="{S.TERTIARY}"/></a:accent3>'
        f'<a:accent4><a:srgbClr val="{S.ACCENT_4}"/></a:accent4>'
        f'<a:accent5><a:srgbClr val="{S.ACCENT_5}"/></a:accent5>'
        f'<a:accent6><a:srgbClr val="{S.TEXT_SUB}"/></a:accent6>'
        f'<a:hlink><a:srgbClr val="{S.PRIMARY}"/></a:hlink>'
        f'<a:folHlink><a:srgbClr val="{S.TEXT_SUB}"/></a:folHlink>'
        '</a:clrScheme>'
    )
    font = (
        '<a:fontScheme name="Noto Sans JP">'
        f'<a:majorFont><a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/><a:cs typeface=""/>'
        f'<a:font script="Jpan" typeface="{FONT}"/></a:majorFont>'
        f'<a:minorFont><a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/><a:cs typeface=""/>'
        f'<a:font script="Jpan" typeface="{FONT}"/></a:minorFont>'
        '</a:fontScheme>'
    )
    import re
    xml = re.sub(r"<a:clrScheme.*?</a:clrScheme>", clr, xml, flags=re.S)
    xml = re.sub(r"<a:fontScheme.*?</a:fontScheme>", font, xml, flags=re.S)
    theme._blob = xml.encode("utf-8")


# ============================================================ ガイド線
def set_guides(prs):
    """表示ガイドを版面境界へ張り直す（既定テンプレート由来の4:3中心線を置換）"""
    import re
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    part = prs.part.part_related_by(RT.VIEW_PROPS)
    xml = part.blob.decode("utf-8")
    guides = "".join(
        f'<p:guide orient="{o}" pos="{S.guide_pos(mm)}"/>' for o, mm, _ in S.GUIDES)
    lst = f"<p:guideLst>{guides}</p:guideLst>"
    if "<p:guideLst>" in xml:
        xml = re.sub(r"<p:guideLst>.*?</p:guideLst>", lst, xml, flags=re.S)
    else:
        xml = xml.replace("</p:cSldViewPr>", lst + "</p:cSldViewPr>")
    # ガイドを表示状態にし、グリッド間隔を余白スケールの基準単位に合わせる
    xml = re.sub(r"<p:cSldViewPr([^>]*?)>",
                 lambda m: "<p:cSldViewPr" + re.sub(r'\s*showGuides="[^"]*"', "", m.group(1))
                           + ' showGuides="1">', xml, count=1)
    xml = re.sub(r'<p:gridSpacing cx="\d+" cy="\d+"/>',
                 f'<p:gridSpacing cx="{S.GRID_EMU}" cy="{S.GRID_EMU}"/>', xml)
    part._blob = xml.encode("utf-8")


# ============================================================ マスター
def build_master(prs):
    master = prs.slide_masters[0]
    cSld = master._element.find(qn("p:cSld"))
    tree = master.shapes._spTree
    clear_shapes(tree)
    set_bg(cSld, S.WHITE)

    tree.append(ph_sp(2, "タイトル", "title", None, S.LEFT, S.TITLE_Y, S.CONTENT_W, S.TITLE_H,
                      S.T["page_title"], S.TEXT, prompt="マスタータイトルの書式設定"))
    tree.append(ph_sp(3, "本文", "body", 1, S.LEFT, S.BODY_TOP, S.CONTENT_W, S.BODY_H,
                      S.T["body"], S.TEXT, prompt="マスターテキストの書式設定"))
    tree.append(ph_sp(4, "フッター", "ftr", 21, S.LEFT, S.FOOT_Y, S.span_w(5), S.FOOT_H,
                      S.T["note"], S.TEXT_SUB, prompt=""))
    tree.append(sldnum_sp(5, S.col_x(5), S.FOOT_Y, S.COL_W, S.FOOT_H, S.T["note"], S.TEXT_SUB))

    # マスターの既定テキストスタイル
    def lvl(st, color, n):
        a, k = rpr(st, color)
        return (f'<a:lvl{n}pPr marL="0" indent="0" algn="l">'
                f'<a:lnSpc><a:spcPts val="{st["lnPts"]}"/></a:lnSpc>'
                f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
                f'<a:defRPr {a}>{k}</a:defRPr></a:lvl{n}pPr>')

    old = master._element.find(qn("p:txStyles"))
    if old is not None:
        master._element.remove(old)
    body_lvls = "".join([
        lvl(S.T["body"], S.TEXT, 1),
        lvl(S.T["body"], S.TEXT, 2),
        lvl(S.T["dense"], S.TEXT_SUB, 3),
        lvl(S.T["dense"], S.TEXT_SUB, 4),
        lvl(S.T["dense"], S.TEXT_SUB, 5),
    ])
    other_lvls = "".join([lvl(S.T["body"], S.TEXT, i) for i in range(1, 6)])
    master._element.append(parse_xml(
        f'<p:txStyles {nsdecls("p", "a")}>'
        f'<p:titleStyle>{lvl(S.T["page_title"], S.TEXT, 1)}</p:titleStyle>'
        f'<p:bodyStyle>{body_lvls}</p:bodyStyle>'
        f'<p:otherStyle>{other_lvls}</p:otherStyle>'
        f'</p:txStyles>'
    ))


# ============================================================ レイアウト
def _header_footer(tree, start_id=2, eyebrow="資料名・章名を入力",
                   title="ページタイトルを入力", rule=True):
    i = start_id
    tree.append(ph_sp(i, "章名", "body", 10, S.LEFT, S.EYEBROW_Y, S.CONTENT_W, S.EYEBROW_H,
                      S.T["note"], S.TEXT_SUB, prompt=eyebrow, anchor="b")); i += 1
    tree.append(ph_sp(i, "ページタイトル", "title", None, S.LEFT, S.TITLE_Y, S.CONTENT_W, S.TITLE_H,
                      S.T["page_title"], S.TEXT, prompt=title)); i += 1
    if rule:
        tree.append(rect_sp(i, "見出し罫", S.LEFT, S.RULE_Y, S.CONTENT_W, S.RULE_H, fill=S.PRIMARY)); i += 1
    tree.append(rect_sp(i, "フッター罫", S.LEFT, S.FRULE_Y, S.CONTENT_W, S.FRULE_H, fill=S.TEXT_SUB)); i += 1
    tree.append(ph_sp(i, "注記", "body", 11, S.LEFT, S.FOOT_Y, S.span_w(5), S.FOOT_H,
                      S.T["note"], S.TEXT_SUB, prompt="注記・出典を入力")); i += 1
    tree.append(sldnum_sp(i, S.col_x(5), S.FOOT_Y, S.COL_W, S.FOOT_H,
                          S.T["note"], S.TEXT_SUB)); i += 1
    return i


def _prepare(layout, name, bg=None):
    layout._element.find(qn("p:cSld")).set("name", name)
    layout._element.set("showMasterSp", "0")
    tree = layout.shapes._spTree
    clear_shapes(tree)
    if bg:
        set_bg(layout._element.find(qn("p:cSld")), bg)
    return tree


def build_layouts(prs):
    master = prs.slide_masters[0]
    layouts = list(master.slide_layouts)
    need = 8
    for lay in layouts[need:]:
        master.slide_layouts.remove(lay)
    L = list(master.slide_layouts)

    # ---- 0. 表紙（白） -------------------------------------------------
    t = _prepare(L[0], "表紙（白）")
    t.append(rect_sp(2, "天面バー", 0, 0, S.PAGE_W_MM, 6.0, fill=S.PRIMARY))
    t.append(ph_sp(3, "資料区分", "body", 10, S.LEFT, 92.0, S.CONTENT_W, 6.0,
                   S.T["cover_meta"], S.TEXT_SUB, prompt="資料区分・シリーズ名", anchor="b"))
    t.append(ph_sp(4, "資料タイトル", "ctrTitle", None, S.LEFT, 100.12, S.CONTENT_W, 36.0,
                   S.T["cover_title"], S.TEXT, prompt="資料タイトルを入力"))
    t.append(rect_sp(5, "タイトル罫", S.LEFT, 144.59, S.span_w(2), 1.0, fill=S.PRIMARY))
    t.append(ph_sp(6, "サブタイトル", "subTitle", 1, S.LEFT, 151.94, S.CONTENT_W, 16.0,
                   S.T["cover_sub"], S.TEXT_SUB, prompt="サブタイトルを入力"))
    t.append(ph_sp(7, "発行者", "body", 11, S.LEFT, 265.67, S.CONTENT_W, 14.4,
                   S.T["cover_meta"], S.TEXT, prompt="組織名・部署名", anchor="b"))

    # ---- 1. 表紙（濃色） -----------------------------------------------
    t = _prepare(L[1], "表紙（濃色）", bg=S.PRIMARY)
    t.append(ph_sp(2, "資料区分", "body", 10, S.LEFT, 92.0, S.CONTENT_W, 6.0,
                   S.T["cover_meta"], S.ACCENT_5, prompt="資料区分・シリーズ名", anchor="b"))
    t.append(ph_sp(3, "資料タイトル", "ctrTitle", None, S.LEFT, 100.12, S.CONTENT_W, 36.0,
                   S.T["cover_title"], S.WHITE, prompt="資料タイトルを入力"))
    t.append(rect_sp(4, "タイトル罫", S.LEFT, 144.59, S.span_w(2), 1.0, fill=S.WHITE))
    t.append(ph_sp(5, "サブタイトル", "subTitle", 1, S.LEFT, 151.94, S.CONTENT_W, 16.0,
                   S.T["cover_sub"], S.ACCENT_5, prompt="サブタイトルを入力"))
    t.append(ph_sp(6, "発行者", "body", 11, S.LEFT, 265.67, S.CONTENT_W, 14.4,
                   S.T["cover_meta"], S.WHITE, prompt="組織名・部署名", anchor="b"))

    # ---- 2. 章扉 -------------------------------------------------------
    t = _prepare(L[2], "章扉", bg=S.BG_TINT)
    t.append(ph_sp(2, "章番号", "body", 10, S.LEFT, 115.0, S.CONTENT_W, 6.5,
                   S.T["section_no"], S.PRIMARY, prompt="SECTION 01", anchor="b"))
    t.append(ph_sp(3, "章タイトル", "title", None, S.LEFT, 123.6, S.CONTENT_W, 13.0,
                   S.T["section_ttl"], S.TEXT, prompt="章タイトルを入力"))
    t.append(rect_sp(4, "章罫", S.LEFT, 145.07, S.span_w(2), 1.0, fill=S.PRIMARY))
    t.append(ph_sp(5, "章概要", "body", 11, S.LEFT, 152.42, S.span_w(4), 30.0,
                   S.T["body"], S.TEXT_SUB, prompt="この章で扱う内容を入力"))

    # ---- 3. 標準（1カラム） --------------------------------------------
    t = _prepare(L[3], "標準（1カラム）")
    i = _header_footer(t)
    t.append(ph_sp(i, "本文", "body", 1, S.LEFT, S.BODY_TOP, S.CONTENT_W, S.BODY_H,
                   S.T["body"], S.TEXT, prompt="本文を入力"))

    # ---- 4. 2カラム ----------------------------------------------------
    t = _prepare(L[4], "2カラム")
    i = _header_footer(t)
    hh, gap = 8.0, S.S2
    for n, (idx_h, idx_b, x) in enumerate([(1, 2, S.LEFT), (3, 4, S.col_x(3))]):
        side = "左" if n == 0 else "右"
        t.append(ph_sp(i, f"{side}見出し", "body", idx_h, x, S.BODY_TOP, S.HALF_W, hh,
                       S.T["h2"], S.TEXT, prompt=f"{side}カラム見出し")); i += 1
        t.append(ph_sp(i, f"{side}本文", "body", idx_b, x, S.BODY_TOP + hh + gap,
                       S.HALF_W, S.BODY_H - hh - gap,
                       S.T["body"], S.TEXT, prompt="本文を入力")); i += 1

    # ---- 5. 図表（フル幅） ---------------------------------------------
    t = _prepare(L[5], "図表（フル幅）")
    i = _header_footer(t)
    cap_h, fig_y = 10.0, S.BODY_TOP + 7.0 + S.S1
    cap_y = S.BODY_BOT - cap_h
    fig_h = cap_y - S.S2 - fig_y
    t.append(ph_sp(i, "図表見出し", "body", 1, S.LEFT, S.BODY_TOP, S.CONTENT_W, 7.0,
                   S.T["h3"], S.TEXT, prompt="図表の見出しを入力")); i += 1
    t.append(rect_sp(i, "図表パネル", S.LEFT, fig_y, S.CONTENT_W, fig_h, fill=S.BG_TINT)); i += 1
    t.append(ph_sp(i, "図表エリア", "body", 2, S.LEFT, fig_y, S.CONTENT_W, fig_h,
                   S.T["dense"], S.TEXT_SUB, align="ctr", anchor="ctr",
                   prompt="グラフ・表・画像を配置")); i += 1
    t.append(ph_sp(i, "図表キャプション", "body", 3, S.LEFT, cap_y, S.CONTENT_W, cap_h,
                   S.T["note"], S.TEXT_SUB, prompt="出典・データの期間・注記を入力")); i += 1

    # ---- 6. 白紙（ヘッダー・フッターのみ） -------------------------------
    t = _prepare(L[6], "白紙（ヘッダー・フッターのみ）")
    _header_footer(t)

    # ---- 7. 裏表紙 -----------------------------------------------------
    t = _prepare(L[7], "裏表紙", bg=S.PRIMARY)
    t.append(ph_sp(2, "組織名", "title", None, S.LEFT, 120.0, S.CONTENT_W, 13.0,
                   S.T["section_ttl"], S.WHITE, prompt="組織名を入力"))
    t.append(rect_sp(3, "裏表紙罫", S.LEFT, 141.5, S.span_w(2), 1.0, fill=S.WHITE))
    t.append(ph_sp(4, "連絡先", "body", 10, S.LEFT, 148.9, S.span_w(4), 40.0,
                   S.T["body"], S.ACCENT_5, prompt="問い合わせ先・URLを入力"))
    t.append(text_sp(5, "出典", S.LEFT, 262.0, S.CONTENT_W, 18.0, [
        para(S.ATTRIBUTION, S.T["note"], S.ACCENT_5),
        para(S.ATTRIBUTION_NOTE, S.T["note"], S.ACCENT_5),
    ]))
    return L


# ============================================================ サンプルスライド
def P(s):
    """para() の出力に名前空間を付けて要素化"""
    return parse_xml(s.replace("<a:p>", '<a:p %s>' % nsdecls("a", "p"), 1))


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def fill_ph(slide, idx, paras):
    shape = ph(slide, idx)
    tx = shape.text_frame._txBody
    for p in tx.findall(qn("a:p")):
        tx.remove(p)
    for s in paras:
        tx.append(P(s))
    return shape


def add_num(slide, n_id=90):
    slide.shapes._spTree.append(
        sldnum_sp(n_id, S.col_x(5), S.FOOT_Y, S.COL_W, S.FOOT_H, S.T["note"], S.TEXT_SUB))


def foot_attr(slide):
    """フッターの注記欄に出典行を入れる。

    B5判では版面幅(143.9mm)より出典行の実寸(約148mm)が長く、1行に収まらない。
    折り返すとフッター領域をはみ出すため、B5では入れない。
    出典は規約ページ（利用上の注意）と裏表紙に集約する
    （docs/design-system/_ATTRIBUTION.md「READMEやNOTICEに集約して記載」の運用）。
    """
    if S.CONTENT_W >= 160.0:
        fill_ph(slide, 11, [para(S.ATTRIBUTION, S.T["note"], S.TEXT_SUB)])


def build_slides(prs, L):
    B, G, P_, W = S.TEXT, S.TEXT_SUB, S.PRIMARY, S.WHITE
    t = S.T

    # ---- 1 表紙 --------------------------------------------------------
    s = prs.slides.add_slide(L[0])
    fill_ph(s, 10, [para("業務報告書テンプレート", t["cover_meta"], G)])
    fill_ph(s, 0, [para("資料タイトルが入ります", t["cover_title"], B)])
    fill_ph(s, 1, [para(f"サブタイトルが入ります（{S.PAPER_NAME}・デジタル庁デザインシステム準拠）",
                        t["cover_sub"], G)])
    fill_ph(s, 11, [para("組織名・部署名", t["cover_meta"], B),
                    para("YYYY年MM月", t["cover_meta"], G)])

    # ---- 2 目次 --------------------------------------------------------
    s = prs.slides.add_slide(L[3]); add_num(s)
    fill_ph(s, 10, [para("本資料について", t["note"], G)])
    fill_ph(s, 0, [para("目次", t["page_title"], B)])
    toc = [("01", "本テンプレートの目的"), ("02", "レイアウトの考え方"),
           ("03", "タイポグラフィとカラー"), ("04", "作図・作表の指針"),
           ("05", "利用上の注意と出典")]
    paras = []
    for i, (no, ttl) in enumerate(toc):
        paras.append(para(f"{no}　{ttl}", t["h2"], B, space_before_pt=0 if i == 0 else 14))
        paras.append(para("章の概要を1〜2行で記述します。", t["body"], G, space_before_pt=3))
    fill_ph(s, 1, paras)
    fill_ph(s, 11, [para("注記・出典を入力", t["note"], G)])

    # ---- 3 章扉 --------------------------------------------------------
    s = prs.slides.add_slide(L[2])
    fill_ph(s, 10, [para("SECTION 02", t["section_no"], P_)])
    fill_ph(s, 0, [para("レイアウトの考え方", t["section_ttl"], B)])
    fill_ph(s, 11, [para(f"マージン・カラム・ガターの定義と、{S.PAPER_NAME}での版面の取り方を示します。",
                         t["body"], G)])

    # ---- 4 標準1カラム --------------------------------------------------
    s = prs.slides.add_slide(L[3]); add_num(s)
    fill_ph(s, 10, [para("02 レイアウトの考え方", t["note"], G)])
    fill_ph(s, 0, [para("版面（グリッド）の定義", t["page_title"], B)])
    body = [
        para(f"{S.PAPER_NAME}（{S.PAGE_W_MM:.0f}×{S.PAGE_H_MM:.0f}mm）の版面は、デジタル庁デザインシステムのレイアウト規定"
             "（マージン・カラム・ガター）に従って構成します。",
             t["body"], B),
        para("グリッドの構成要素", t["h3"], B, space_before_pt=14),
        para(f"マージン：左右 {mm1(S.MARGIN_X)}mm／上下 {mm1(S.MARGIN_T)}mm。版面幅は {mm1(S.CONTENT_W)}mm となります。",
             t["body"], B, bullet=True, space_before_pt=5),
        para(f"カラム：{S.COLS}カラム。1カラムの幅は {mm1(S.COL_W)}mm"
             f"（本文文字サイズ{S.BASE_PX}pxの{S.COL_MULT}倍）です。",
             t["body"], B, bullet=True, space_before_pt=3),
        para(f"ガター：{mm1(S.GUTTER)}mm。本文文字サイズの2倍を確保し、隣接カラムの誤読を防ぎます。",
             t["body"], B, bullet=True, space_before_pt=3),
        para("余白スケール", t["h3"], B, space_before_pt=14),
        para(f"基準単位は 8 CSS px（{mm1(S.U)}mm）。1・2・3・4・8倍の5段階に絞って使用します。",
             t["body"], B, space_before_pt=5),
        para(f"{mm1(S.S1)}mm／{mm1(S.S2)}mm／{mm1(S.S3)}mm／{mm1(S.S4)}mm／{mm1(S.S8)}mm",
             t["dense_b"], P_, space_before_pt=5),
        para("要素の関係が近いほど小さい余白を、階層が変わるところには大きい余白を与えます。",
             t["body"], B, space_before_pt=5),
    ]
    fill_ph(s, 1, body)
    foot_attr(s)

    # ---- 4b グリッド図 ----------------------------------------------------
    s = prs.slides.add_slide(L[6]); add_num(s)
    fill_ph(s, 10, [para("02 レイアウトの考え方", t["note"], G)])
    fill_ph(s, 0, [para("グリッド図", t["page_title"], B)])
    foot_attr(s)
    sid = 500
    # B5は版面が狭く凡例の値が折り返すため、図を1カラム分ゆずる（A4は従来どおり 4:2）
    _dia_cols = 4 if S.PROFILE == "a4" else 3
    _leg_cols = S.COLS - _dia_cols
    dia_w = S.span_w(_dia_cols)
    k = dia_w / S.PAGE_W_MM
    ox, oy = S.LEFT, S.BODY_TOP + 4.0
    def dx(v): return ox + v * k
    def dy(v): return oy + v * k
    def dk(v): return v * k
    # 用紙
    s.shapes._spTree.append(rect_sp(sid, "用紙", ox, oy, dk(S.PAGE_W_MM), dk(S.PAGE_H_MM),
                                    fill=S.WHITE, line=S.TEXT_SUB, line_w_mm=0.2)); sid += 1
    # ヘッダー・フッター帯
    s.shapes._spTree.append(rect_sp(sid, "ヘッダー帯", dx(S.LEFT), dy(S.EYEBROW_Y),
                                    dk(S.CONTENT_W), dk(S.BODY_TOP - S.EYEBROW_Y),
                                    fill=S.BG_TINT)); sid += 1
    s.shapes._spTree.append(rect_sp(sid, "フッター帯", dx(S.LEFT), dy(S.FRULE_Y),
                                    dk(S.CONTENT_W), dk(S.FOOT_BOT - S.FRULE_Y),
                                    fill=S.BG_TINT)); sid += 1
    s.shapes._spTree.append(rect_sp(sid, "見出し罫図", dx(S.LEFT), dy(S.RULE_Y),
                                    dk(S.CONTENT_W), max(dk(S.RULE_H), 0.3),
                                    fill=S.PRIMARY)); sid += 1
    # 6カラム
    for i in range(S.COLS):
        s.shapes._spTree.append(rect_sp(sid, f"カラム{i}", dx(S.col_x(i)), dy(S.BODY_TOP),
                                        dk(S.COL_W), dk(S.BODY_H), fill=S.ACCENT_5)); sid += 1
    # 寸法の引き出し
    s.shapes._spTree.append(text_sp(sid, "図注1", ox, oy + dk(S.PAGE_H_MM) + 3.0,
                                    dia_w, 12.0,
                                    [para("水色＝カラム／淡灰＝ヘッダー・フッター領域",
                                          t["note"], G)])); sid += 1
    # 凡例
    lx = S.col_x(_dia_cols)
    lw = S.span_w(_leg_cols)
    legend = [("用紙", f"{S.PAPER_NAME} {S.PAGE_W_MM:.0f} × {S.PAGE_H_MM:.0f} mm"),
              ("マージン", f"左右 {mm1(S.MARGIN_X)} mm／上下 {mm1(S.MARGIN_T)} mm"),
              ("版面", f"{mm1(S.CONTENT_W)} × {mm1(S.BODY_H)} mm"),
              ("カラム", f"{S.COLS} カラム × {mm1(S.COL_W)} mm"),
              ("ガター", f"{mm1(S.GUTTER)} mm（本文の2倍）"),
              ("ヘッダー領域", f"{mm1(S.MARGIN_T)} 〜 {mm1(S.BODY_TOP)} mm"),
              ("フッター領域", f"{mm1(S.FRULE_Y)} 〜 {mm1(S.FOOT_BOT)} mm")]
    ly = S.BODY_TOP + 4.0
    s.shapes._spTree.append(text_sp(sid, "凡例見出し", lx, ly, lw, 8.0,
                                    [para("寸法", t["h3"], B)])); sid += 1
    ly += 9.0
    for name, val in legend:
        s.shapes._spTree.append(rect_sp(sid, f"凡例罫{name}", lx, ly, lw, 0.15,
                                        fill=S.ACCENT_4)); sid += 1
        s.shapes._spTree.append(text_sp(sid, f"凡例{name}", lx, ly + 2.0, lw, 12.0, [
            para(name, t["dense_b"], B),
            para(val, t["dense"], G, space_before_pt=1)])); sid += 1
        ly += 15.0
    s.shapes._spTree.append(text_sp(sid, "図解説", ox, oy + dk(S.PAGE_H_MM) + 16.0,
                                    dia_w, 50.0, [
        para("使い方", t["h3"], B),
        para("図・表・カードはカラムの左端と右端にそろえます。カラムをまたぐ場合は"
             f"ガター分を含めた幅（2カラム＝{mm1(S.span_w(2))}mm、3カラム＝{mm1(S.span_w(3))}mm、"
             f"{S.COLS}カラム＝{mm1(S.CONTENT_W)}mm）を使います。",
             t["body"], B, space_before_pt=5),
    ]))

    # ---- 5 2カラム ------------------------------------------------------
    s = prs.slides.add_slide(L[4]); add_num(s)
    fill_ph(s, 10, [para("02 レイアウトの考え方", t["note"], G)])
    fill_ph(s, 0, [para("2カラムの使い分け", t["page_title"], B)])
    fill_ph(s, 1, [para("向いている内容", t["h2"], B)])
    fill_ph(s, 2, [
        para("対になる情報（現状と課題、施策と効果など）を並べて比較する場合に使用します。",
             t["body"], B),
        para(f"片側 {mm1(S.HALF_W)}mm（3カラム分）。1行あたりの文字数は本文{S.T['body']['pt']:g}ptで"
             f"約{int(S.HALF_W / (S.T['body']['pt'] * 25.4 / 72))}文字となり、"
             "視線の折り返しが短く読み進めやすい行長です。", t["body"], B, space_before_pt=8),
        para("左右の高さを揃える", t["h3"], B, space_before_pt=14),
        para("上端を必ず揃え、下端は無理に揃えません。", t["body"], B, bullet=True, space_before_pt=5),
        para("片側が極端に短い場合は1カラムに切り替えます。", t["body"], B, bullet=True, space_before_pt=3),
    ])
    fill_ph(s, 3, [para("向いていない内容", t["h2"], B)])
    fill_ph(s, 4, [
        para("手順や時系列など、順序に意味がある内容は2カラムに分けません。"
             "視線が左右に往復し、読み順が壊れるためです。", t["body"], B),
        para("読み上げ順序と表示順序を一致させることは、デジタル庁デザインシステムが"
             "レイアウトのアクセシビリティ要件として挙げている事項です。",
             t["body"], B, space_before_pt=8),
        para("長い表や横長のグラフも、2カラムではなく「図表（フル幅）」レイアウトを使用します。",
             t["body"], B, space_before_pt=8),
    ])
    foot_attr(s)

    # ---- 6 サマリー（KPIカード） ------------------------------------------
    s = prs.slides.add_slide(L[6]); add_num(s)
    fill_ph(s, 10, [para("01 本テンプレートの目的", t["note"], G)])
    fill_ph(s, 0, [para("サマリー数値の見せ方", t["page_title"], B)])
    fill_ph(s, 11, [para("※ 数値はサンプルです。", t["note"], G)])
    sid = 100
    card_w, card_h = S.HALF_W, 34.0
    cards = [("登録利用者数", "128,450", "人", "前月比 +3.2%"),
             ("平均処理日数", "4.2", "日", "前月比 -0.8日"),
             ("オンライン申請率", "76.4", "%", "前月比 +1.5pt"),
             ("問い合わせ件数", "1,032", "件", "前月比 -12.4%")]
    for n, (label, val, unit, sub) in enumerate(cards):
        x = S.LEFT if n % 2 == 0 else S.col_x(3)
        y = S.BODY_TOP + (n // 2) * (card_h + S.S4)
        s.shapes._spTree.append(rect_sp(sid, f"カード{n}", x, y, card_w, card_h, fill=S.BG_TINT)); sid += 1
        s.shapes._spTree.append(rect_sp(sid, f"カード罫{n}", x, y, 1.2, card_h, fill=S.PRIMARY)); sid += 1
        s.shapes._spTree.append(text_sp(sid, f"カード文言{n}", x + S.S3, y + S.S3,
                                        card_w - S.S3 * 2, card_h - S.S3 * 2, [
            para(label, t["label"], G),
            para(f"{val} {unit}", t["kpi_num"], P_, space_before_pt=6),
            para(sub, t["note"], G, space_before_pt=4),
        ])); sid += 1
    y2 = S.BODY_TOP + 2 * (card_h + S.S4)
    s.shapes._spTree.append(text_sp(sid, "解説", S.LEFT, y2, S.CONTENT_W, 60.0, [
        para("カードの設計", t["h3"], B),
        para(f"ラベル（{S.T['label']['pt']:g}pt）→ 数値（{S.T['kpi_num']['pt']:g}pt）→ "
             f"補足（{S.T['note']['pt']:g}pt）の3階層で、視線が数値に"
             f"止まるようにしています。カードの内側パディングは余白スケールの {mm1(S.S3)}mm、"
             f"カード間のガターは {mm1(S.GUTTER)}mm です。", t["body"], B, space_before_pt=5),
        para("色だけで増減を示さない", t["h3"], B, space_before_pt=12),
        para("増減は「+3.2%」「-0.8日」のように符号と単位を文字で明記します。"
             "色の違いだけに意味を持たせると、色覚特性によって情報が伝わりません。",
             t["body"], B, space_before_pt=5),
    ]))

    # ---- 7 図表（フル幅） -------------------------------------------------
    s = prs.slides.add_slide(L[5]); add_num(s)
    fill_ph(s, 10, [para("04 作図・作表の指針", t["note"], G)])
    fill_ph(s, 0, [para("グラフの配置", t["page_title"], B)])
    fill_ph(s, 1, [para("図1　月次推移（サンプル）", t["h3"], B)])
    fill_ph(s, 2, [para(f"ここにグラフ・表・画像を配置します（版面幅 {mm1(S.CONTENT_W)}mm）",
                        t["dense"], G, align="ctr")])
    fill_ph(s, 3, [para("出典：〇〇調査（YYYY年MM月実施）。n=1,000。四捨五入のため合計が"
                        "100%にならない場合があります。", t["note"], G)])
    foot_attr(s)

    # ---- 8 表 ------------------------------------------------------------
    s = prs.slides.add_slide(L[6]); add_num(s)
    fill_ph(s, 10, [para("04 作図・作表の指針", t["note"], G)])
    fill_ph(s, 0, [para("表の体裁", t["page_title"], B)])
    fill_ph(s, 11, [para("※ 数値はサンプルです。", t["note"], G)])
    rows = [["項目", "前年度", "本年度", "増減"],
            ["オンライン申請件数", "82,140", "104,880", "+27.7%"],
            ["窓口来訪件数", "41,320", "33,015", "-20.1%"],
            ["平均処理日数", "5.8 日", "4.2 日", "-1.6 日"],
            ["利用者満足度", "68.2 %", "74.9 %", "+6.7 pt"]]
    add_table(s, rows, S.LEFT, S.BODY_TOP, S.CONTENT_W,
              [S.span_w(3) - S.GUTTER, 0, 0, 0])
    ty = S.BODY_TOP + 12.0 + 4 * 10.0 + S.S4
    s.shapes._spTree.append(text_sp(200, "表解説", S.LEFT, ty, S.CONTENT_W, 70.0, [
        para("表1　主要指標の前年度比較", t["note"], G),
        para("罫線と塗りの使い分け", t["h3"], B, space_before_pt=14),
        para(f"縦罫は引かず、行の区切りだけを {mm1(S.FRULE_H)}mm の横罫で示します。"
             "罫線はデジタル庁デザインシステムの非テキスト要素の規定に従い、"
             "背景とのコントラスト比 3:1 以上を確保した色を使用しています。",
             t["body"], B, space_before_pt=5),
        para("数値は右揃え、桁を揃える", t["h3"], B, space_before_pt=12),
        para("数値列は右揃えにして桁位置を合わせます。単位はヘッダーまたは各セルに明記し、"
             "単位のない裸の数値を並べません。", t["body"], B, space_before_pt=5),
    ]))

    # ---- 9 タイポグラフィ ------------------------------------------------
    s = prs.slides.add_slide(L[6]); add_num(s)
    fill_ph(s, 10, [para("03 タイポグラフィとカラー", t["note"], G)])
    fill_ph(s, 0, [para("書体とサイズ", t["page_title"], B)])
    foot_attr(s)
    sid, y = 300, S.BODY_TOP
    _k0 = S.CONTENT_W / A4_CONTENT_W
    s.shapes._spTree.append(text_sp(sid, "書体説明", S.LEFT, y, S.CONTENT_W, 16.0, [
        para("書体は Noto Sans JP を使用します。サイズはデジタル庁デザインシステムの"
             "テキストスタイル（CSS px）を 0.75 倍して pt に換算した値です。",
             t["body"], B)])); sid += 1
    y += 18.0 * _k0
    # 見出しの行高は判型で変わる。実際の値域から文言を組み立てる（A4では "150%"）
    _hs = sorted({t[k]["lh"] for k in ("section_ttl", "page_title", "h2")})
    _head_lh = f"{_hs[0]}%" if len(_hs) == 1 else f"{_hs[0]}〜{_hs[-1]}%"
    _ratio = (f"本文{t['body']['lh']}%／見出し{_head_lh}／表・注記{t['dense']['lh']}%")
    if S.PROFILE == "a4":
        _lh_text = (f"行間はフォントサイズに対する比率で固定しています（{_ratio}）。"
                    "PowerPoint上は「固定値（pt）」で指定しているため、"
                    "書体を入れ替えても行の位置がずれません。")
        _spc_text = ("字間はデジタル庁デザインシステムの指定（0％・1％・2％）を pt に換算して"
                     "設定済みです。個別に変更しないでください。")
    else:
        # B5は本文高がA4より約36mm少ないため、同じ内容を短くまとめる
        _lh_text = (f"行間は比率で固定しています（{_ratio}）。"
                    "PowerPoint上は「固定値（pt）」指定のため、書体を替えても行位置がずれません。")
        _spc_text = "字間はデジタル庁デザインシステムの指定（0％・1％・2％）を pt に換算済みです。"
    # 見本の文字列は版面が狭いほど短くする（判型をまたいで折り返さないように）
    long_s = "見本 Sample 0123" if S.PROFILE == "a4" else "見本 Sample"
    top_s = "見本 Sample" if S.PROFILE == "a4" else "見本 Abc"
    scale = [("資料タイトル", "cover_title", top_s),
             ("章タイトル", "section_ttl", long_s),
             ("ページタイトル", "page_title", long_s),
             ("見出し2", "h2", long_s),
             ("見出し3", "h3", "見本 Sample 0123"),
             ("本文", "body", "見本 Sample 0123"),
             ("注記・表", "dense", "見本 Sample 0123")]
    # 欄幅は版面幅に比例させる（A4では 30.0 / 46.0 と同値）
    _k = S.CONTENT_W / A4_CONTENT_W
    # 用途ラベルは最長 "ページタイトル"(7文字) が1行に収まる幅を確保する
    # 行送りはB5では半分に詰める（本文高がA4より約36mm少ないため）
    _row_gap = S.S1 if S.PROFILE == "a4" else S.S1 * 0.5
    _label_min = 7 * t["dense"]["pt"] * 25.4 / 72 + 1.2
    name_w, meta_w = max(30.0 * _k, _label_min), 46.0 * _k
    spec_x = S.LEFT + name_w + 4.0
    spec_w = S.RIGHT - meta_w - 2.0 - spec_x
    for name, key, sample in scale:
        st = t[key]
        h = st["lnPts"] / 100 * 25.4 / 72 + 4.0
        s.shapes._spTree.append(text_sp(sid, f"用途{key}", S.LEFT, y, name_w, h,
                                        [para(name, t["dense"], G)], anchor="ctr")); sid += 1
        s.shapes._spTree.append(text_sp(sid, f"見本{key}", spec_x, y, spec_w, h,
                                        [para(sample, st, B)], anchor="ctr")); sid += 1
        s.shapes._spTree.append(text_sp(sid, f"仕様{key}", S.RIGHT - meta_w, y, meta_w, h,
                                        [para(f'{dads_token(st)}{"" if dads_defined(st) else "※"}'
                                              f'／{st["pt"]:g}pt', t["dense"], P_,
                                              align="r")], anchor="ctr")); sid += 1
        y += h + _row_gap
        s.shapes._spTree.append(rect_sp(sid, f"区切{key}", S.LEFT, y - _row_gap / 2,
                                        S.CONTENT_W, 0.15, fill=S.ACCENT_4)); sid += 1
    if any(not dads_defined(t[k]) for _, k, _ in scale):
        s.shapes._spTree.append(text_sp(sid, "トークン注記", S.LEFT, y, S.CONTENT_W, 9.0, [
            para("※ はデジタル庁デザインシステムのテキストスタイル表に無い組み合わせです"
                 "（読みやすさのため行間を広げています）。", t["note"], G)])); sid += 1
        y += 9.0
    y += S.S3
    s.shapes._spTree.append(text_sp(sid, "行間説明", S.LEFT, y, S.CONTENT_W, 70.0, [
        para("行間と字間", t["h2"], B),
        para(_lh_text, t["body"], B, space_before_pt=6),
        para(_spc_text, t["body"], B, space_before_pt=6),
    ] + ([
        para("DADSの14 CSS px 未満のサイズは使用しません。",
             t["body"], B, bullet=True, space_before_pt=8),
        para("強調は太さレベル（Bold）で行い、下線や斜体は使用しません。",
             t["body"], B, bullet=True, space_before_pt=3),
    ] if S.PROFILE == "a4" else [
        para("DADSの14 CSS px 未満は使用しません。強調は Bold で行い、下線・斜体は使いません。",
             t["body"], B, bullet=True, space_before_pt=8),
    ])))

    # ---- 10 カラー ---------------------------------------------------------
    s = prs.slides.add_slide(L[6]); add_num(s)
    fill_ph(s, 10, [para("03 タイポグラフィとカラー", t["note"], G)])
    fill_ph(s, 0, [para("カラーと使用可否", t["page_title"], B)])
    foot_attr(s)
    sid, y = 400, S.BODY_TOP
    s.shapes._spTree.append(text_sp(sid, "色説明", S.LEFT, y, S.CONTENT_W, 16.0, [
        para("デジタル庁「ダッシュボードイメージ作成ツールキット」のテーマ配色を採用し、"
             "白背景に対するコントラスト比を実測して用途を決めています。",
             t["body"], B)])); sid += 1
    y += 18.0
    swatches = [
        (S.PRIMARY,   "プライマリー",     "#0C21BA", "白背景 10.81:1", "テキスト・罫線・面"),
        (S.SECONDARY, "セカンダリー",     "#2E4EE7", "白背景 6.25:1",  "テキスト・罫線・面"),
        (S.TERTIARY,  "ターシャリー",     "#4F7AE9", "白背景 3.97:1",  "罫線・面のみ（文字には使えません）"),
        (S.TEXT_SUB,  "補助テキスト・罫", "#626264", "白背景 6.09:1",  "補助文字・罫線"),
        (S.ACCENT_4,  "淡アクセント",     "#99B0EC", "白背景 2.15:1",  "面の塗りのみ"),
        (S.ACCENT_5,  "最淡アクセント",   "#CFDCF0", "白背景 1.39:1",  "面の塗りのみ"),
    ]
    cw, pitch = S.HALF_W, 34.0
    for n, (hexv, nm, code, ratio, use) in enumerate(swatches):
        x = S.LEFT if n % 2 == 0 else S.col_x(3)
        yy = y + (n // 2) * pitch
        s.shapes._spTree.append(rect_sp(sid, f"色{n}", x, yy, cw, 12.0, fill=hexv)); sid += 1
        s.shapes._spTree.append(text_sp(sid, f"色名{n}", x, yy + 13.5, cw, 16.0, [
            para(nm, t["dense_b"], B),
            para(f"{code}　{ratio}", t["dense"], G, space_before_pt=1),
            para(use, t["dense"], G, space_before_pt=1),
        ])); sid += 1
    y += 3 * pitch + S.S2
    s.shapes._spTree.append(text_sp(sid, "色の原則", S.LEFT, y, S.CONTENT_W, 90.0, [
        para("色の使用原則", t["h2"], B),
    ] + ([
        para("テキストと背景のコントラスト比は常に 4.5:1 以上、罫線などの非テキスト要素は "
             "3:1 以上を確保します。上の比率は実測値です。",
             t["body"], B, bullet=True, space_before_pt=6),
        para("コントラスト比が 4.5:1 未満の色を文字に使わないでください。"
             "ターシャリー以下の3色は面や罫線専用です。",
             t["body"], B, bullet=True, space_before_pt=3),
        para("色だけで情報を区別せず、文字・記号・形状を併用します。",
             t["body"], B, bullet=True, space_before_pt=3),
        para("フォーカスインジケーターの配色（Yellow-300 と Black の2重構造）は"
             "いかなる場合も変更してはいけません。",
             t["body"], B, bullet=True, space_before_pt=3),
        para("背景に淡色（#F7F8FB）を敷いた場合もコントラスト比を再確認してください"
             "（補助グレーは 5.73:1、黒は 19.77:1）。",
             t["body"], B, bullet=True, space_before_pt=3),
    ] if S.PROFILE == "a4" else [
        # B5は本文高がA4より約36mm少ないため、同じ内容を短くまとめる
        para("テキストは 4.5:1 以上、非テキスト要素は 3:1 以上を確保します"
             "（淡色背景 #F7F8FB でも同様。上の比率は実測値）。",
             t["body"], B, bullet=True, space_before_pt=6),
        para("4.5:1 未満の色は文字に使えません。ターシャリー以下の3色は面・罫線専用です。",
             t["body"], B, bullet=True, space_before_pt=3),
        para("色だけで情報を区別せず、文字・記号・形状を併用します。",
             t["body"], B, bullet=True, space_before_pt=3),
        para("フォーカスインジケーターの配色（Yellow-300 と Black の2重構造）は変更禁止です。",
             t["body"], B, bullet=True, space_before_pt=3),
    ])))

    # ---- 10 利用上の注意 ---------------------------------------------------
    s = prs.slides.add_slide(L[3]); add_num(s)
    fill_ph(s, 10, [para("05 利用上の注意と出典", t["note"], G)])
    fill_ph(s, 0, [para("利用上の注意と出典", t["page_title"], B)])
    fill_ph(s, 1, [
        para("本テンプレートについて", t["h3"], B),
        para("デジタル庁「ダッシュボードイメージ作成ツールキット」(16:9) のレイアウト思想"
             "（マージン・グリッド・ヘッダー／フッターの構成）を踏まえ、"
             f"デジタル庁デザインシステムの基本デザインに従って{S.PAPER_NAME}用に再設計したものです。",
             t["body"], B, space_before_pt=5),
        para("遵守している基準", t["h3"], B, space_before_pt=14),
        para("テキストと背景のコントラスト比 4.5:1 以上／非テキスト要素 3:1 以上",
             t["body"], B, bullet=True, space_before_pt=5),
        para("余白は基準単位 8 CSS px の倍率スケール（5段階）",
             t["body"], B, bullet=True, space_before_pt=3),
        para("カラムのガターは本文文字サイズの2倍",
             t["body"], B, bullet=True, space_before_pt=3),
        para("書体は Noto Sans JP（未導入環境では代替書体に置き換わります）",
             t["body"], B, bullet=True, space_before_pt=3),
        para("出典表記", t["h3"], B, space_before_pt=14),
        para(S.ATTRIBUTION, t["body"], B, space_before_pt=5),
        para(S.ATTRIBUTION_NOTE, t["body"], B, space_before_pt=3),
        para("デジタル庁が作成したかのような態様での公表・利用は禁止されています。"
             "本テンプレートを用いた成果物には、上記の出典行を必ず記載してください。",
             t["body"], B, space_before_pt=3),
    ])
    fill_ph(s, 11, [para("注記・出典を入力", t["note"], G)])

    # ---- 11 裏表紙 ---------------------------------------------------------
    s = prs.slides.add_slide(L[7])
    fill_ph(s, 0, [para("組織名が入ります", t["section_ttl"], W)])
    fill_ph(s, 10, [para("〒000-0000　都道府県市区町村0-0-0", t["body"], S.ACCENT_5),
                    para("https://example.jp/", t["body"], S.ACCENT_5, space_before_pt=3)])


def add_table(slide, rows, x, y, w, col_w):
    """罫線最小の表（ヘッダー塗り＋横罫のみ）"""
    n_r, n_c = len(rows), len(rows[0])
    head_h, row_h = 12.0, 10.0
    free = w - col_w[0]
    widths = [col_w[0]] + [free / (n_c - 1)] * (n_c - 1)
    gf = slide.shapes.add_table(n_r, n_c, Emu(emu(x)), Emu(emu(y)), Emu(emu(w)),
                                Emu(emu(head_h + row_h * (n_r - 1))))
    tbl = gf.table
    # 既定のテーブルスタイル（縞模様）を無効化
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for tag in ("a:tableStyleId",):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(emu(cw))
    tbl.rows[0].height = Emu(emu(head_h))
    for r in range(1, n_r):
        tbl.rows[r].height = Emu(emu(row_h))

    for r in range(n_r):
        for c in range(n_c):
            cell = tbl.cell(r, c)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for el in list(tcPr):
                tcPr.remove(el)
            tcPr.set("marL", str(emu(2.0)))
            tcPr.set("marR", str(emu(2.0)))
            tcPr.set("marT", str(emu(1.5)))
            tcPr.set("marB", str(emu(1.5)))
            tcPr.set("anchor", "ctr")
            # 横罫のみ（上辺は見出し行の上＝太罫、以降は細罫）
            lnB = (f'<a:lnB w="{emu(0.2)}" cap="flat"><a:solidFill>'
                   f'<a:srgbClr val="{S.TEXT_SUB}"/></a:solidFill></a:lnB>')
            lnT = ""
            if r == 0:
                lnT = (f'<a:lnT w="{emu(0.5)}" cap="flat"><a:solidFill>'
                       f'<a:srgbClr val="{S.PRIMARY}"/></a:solidFill></a:lnT>')
            fill_xml = (f'<a:solidFill><a:srgbClr val="{S.BG_TINT}"/></a:solidFill>'
                        if r == 0 else '<a:noFill/>')
            frag = parse_xml(f'<a:x {nsdecls("a")}>{lnT}{lnB}{fill_xml}</a:x>')
            for el in frag:
                tcPr.append(el)

            st = S.T["dense_b"] if r == 0 else S.T["dense"]
            color = S.TEXT if r == 0 else S.TEXT
            align = "l" if c == 0 else "r"
            tx = cell.text_frame._txBody
            for p in tx.findall(qn("a:p")):
                tx.remove(p)
            tx.append(P(para(rows[r][c], st, color, align=align)))
    return gf


# ============================================================ main
def main():
    prs = Presentation()
    prs.slide_width = Emu(emu(S.PAGE_W_MM))
    prs.slide_height = Emu(emu(S.PAGE_H_MM))
    # 既定テンプレート由来の type="screen4x3" は実寸と矛盾するため custom にする
    prs._element.find(qn("p:sldSz")).set("type", "custom")
    rewrite_theme(prs)
    set_guides(prs)
    build_master(prs)
    L = build_layouts(prs)
    build_slides(prs, L)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    freeze_zip_dates(OUT)
    print(f"saved: {OUT}")
    print(f"slide size: {prs.slide_width} x {prs.slide_height} EMU "
          f"({prs.slide_width/S.MM:.1f} x {prs.slide_height/S.MM:.1f} mm)")
    print("layouts:", [l.name for l in prs.slide_masters[0].slide_layouts])
    print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
    print("guides:")
    for o, mm, note in S.GUIDES:
        print(f"   {o:<4} {mm:7.3f}mm  pos={S.guide_pos(mm):<5} {note}")


if __name__ == "__main__":
    main()
